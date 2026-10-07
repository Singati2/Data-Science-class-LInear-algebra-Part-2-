"""Build the Week 2 class slides (Solving Systems of Linear Equations).

Run:  python3 build_slides.py
Out:  Linear_Algebra_Week2_Slides.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------- palette ----
NAVY = "14213D"
TEAL = "1F6F8B"
AMBER = "C97B2B"
AMBER_L = "F2B263"
WHITE = "FFFFFF"
LIGHT = "BFD8D5"
GREY = "5B6675"
TINT = "F1F3F6"
CELL = "E3EAF6"
PIV = "FFE2B8"
RES = "CFE8EC"
RED = "B23A3A"

HEAD = "Cambria"
BODY = "Calibri"

W, H = 13.333, 7.5
FOOT = "Math for Data Science  ·  Linear Algebra, Week 2"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- helpers ----
def text(slide, x, y, w, h, content, size=18, color=NAVY, font=BODY, bold=False,
         italic=False, align="l", anchor="t", spacing=None, after=0, margin=0, wrap=True):
    """content: str (\\n = new paragraph) | list of paragraphs; a paragraph is a str
    or a list of runs (text, {bold, italic, color, size, font})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(content, str):
        paras = content.split("\n")
    elif content and isinstance(content[0], tuple):
        paras = [content]
    else:
        paras = content
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        if after:
            p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else para
        for rt, st in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.italic = st.get("italic", italic)
            f.name = st.get("font", font)
            f.color.rgb = rgb(st.get("color", color))
    return tb


def rect(slide, x, y, w, h, fill=TINT, line=None, lw=1.0, rounded=False, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=NAVY, lw=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    return c


def arrow(slide, x, y, w=0.7, h=0.34, color=TEAL, label=None, lsize=13, lw_=2.2):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    if label:
        text(slide, x - (lw_ - w) / 2, y + h + 0.06, lw_, 0.8, label, size=lsize, color=GREY, align="c")
    return shp


def matrix(slide, x, y, rows, cw=0.58, ch=0.5, size=18, fill=CELL, hl=None, aug=None,
           gap=0.05, label=None, lsize=14, lcolor=GREY, bracket=True, tcolor=NAVY, font=BODY, bcolor=NAVY):
    """Draw a matrix of cells at (x, y). hl = {(r, c): color}. aug = index of the
    last coefficient column (a bar is drawn after it). Returns (w, h)."""
    hl = hl or {}
    n, m = len(rows), len(rows[0])
    bx = 0.14 if bracket else 0.0
    extra = 0.16 if aug is not None else 0.0
    tw = bx * 2 + m * cw + (m - 1) * gap + extra
    th = n * ch + (n - 1) * gap
    for r in range(n):
        cx = x + bx
        for c in range(m):
            if aug is not None and c == aug + 1:
                cx += extra
            col = hl.get((r, c), fill)
            cell = rect(slide, cx, y + r * (ch + gap), cw, ch, fill=col)
            tf = cell.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            rn = p.add_run()
            rn.text = str(rows[r][c])
            rn.font.size = Pt(size)
            rn.font.name = font
            rn.font.color.rgb = rgb(tcolor)
            cx += cw + gap
    if aug is not None:
        ax = x + bx + (aug + 1) * cw + aug * gap + gap / 2 + extra / 2
        line(slide, ax, y - 0.04, ax, y + th + 0.04, color=NAVY, lw=2.5)
    if bracket:
        for kind, bxp in ((MSO_SHAPE.LEFT_BRACKET, x), (MSO_SHAPE.RIGHT_BRACKET, x + tw - bx)):
            b = slide.shapes.add_shape(kind, Inches(bxp), Inches(y - 0.04), Inches(bx), Inches(th + 0.08))
            b.shadow.inherit = False
            b.fill.background()
            b.line.color.rgb = rgb(bcolor)
            b.line.width = Pt(1.75)
    if label:
        text(slide, x - 0.6, y + th + 0.12, tw + 1.2, 0.6, label, size=lsize, color=lcolor, align="c")
    return tw, th


def eqs(slide, x, y, lines, size=22, gap=0.46, font=HEAD, color=NAVY, w=4.2, title=None, tsize=14):
    """A column of equations (strings or run-lists). Returns the bottom y."""
    yy = y
    if title:
        text(slide, x, yy, w, 0.35, title, size=tsize, color=GREY, bold=True)
        yy += 0.38
    for ln in lines:
        text(slide, x, yy, w, gap, ln, size=size, font=font, color=color, anchor="m")
        yy += gap
    return yy


def header(slide, title, eyebrow=None, dark=False):
    col = WHITE if dark else NAVY
    ey = AMBER_L if dark else TEAL
    if eyebrow:
        text(slide, 0.6, 0.42, 10, 0.3, eyebrow.upper(), size=12, color=ey, bold=True)
    text(slide, 0.6, 0.68, 12.1, 0.8, title, size=34, font=HEAD, color=col, bold=True, anchor="t")


def footer(slide, n, dark=False):
    col = LIGHT if dark else GREY
    text(slide, 0.6, 7.02, 8, 0.3, FOOT, size=10, color=col)
    text(slide, 11.9, 7.02, 0.83, 0.3, str(n), size=10, color=col, align="r")


def new_slide(title=None, eyebrow=None, dark=False, notes=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    if title:
        header(s, title, eyebrow, dark)
    footer(s, len(prs.slides), dark)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def defbox(slide, x, y, w, h, body, label="Definition", size=17, fill=TINT, lcolor=TEAL):
    rect(slide, x, y, w, h, fill=fill, rounded=True, radius=0.04)
    text(slide, x + 0.3, y + 0.22, w - 0.6, 0.3, label.upper(), size=12, color=lcolor, bold=True)
    text(slide, x + 0.3, y + 0.56, w - 0.6, h - 0.7, body, size=size, color=NAVY, spacing=1.15, after=6)


def card(slide, x, y, w, h, title, body, fill=TINT, tsize=18, bsize=15, number=None, tcolor=NAVY):
    rect(slide, x, y, w, h, fill=fill, rounded=True, radius=0.04)
    ty = y + 0.22
    if number is not None:
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.25), Inches(y + 0.22), Inches(0.42), Inches(0.42))
        circ.shadow.inherit = False
        circ.fill.solid(); circ.fill.fore_color.rgb = rgb(TEAL); circ.line.fill.background()
        tf = circ.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(number); r.font.size = Pt(14); r.font.bold = True
        r.font.color.rgb = rgb(WHITE); r.font.name = BODY
        text(slide, x + 0.8, ty, w - 1.05, 0.45, title, size=tsize, bold=True, color=tcolor, anchor="m")
        ty += 0.6
    else:
        text(slide, x + 0.25, ty, w - 0.5, 0.4, title, size=tsize, bold=True, color=tcolor)
        ty += 0.5
    text(slide, x + 0.25, ty, w - 0.5, h - (ty - y) - 0.2, body, size=bsize, color=GREY, spacing=1.12, after=4)


def divider(part, title, blurb):
    s = new_slide(dark=True)
    text(s, 0.6, 2.2, 6, 0.4, f"PART {part}", size=16, color=AMBER_L, bold=True)
    text(s, 0.6, 2.7, 11.5, 1.6, title, size=54, font=HEAD, color=WHITE, bold=True)
    text(s, 0.6, 4.5, 9, 1.2, blurb, size=20, color=LIGHT, spacing=1.2)
    return s


def quiz(n, prompt, body_lines, hint=None):
    s = new_slide(dark=True)
    text(s, 0.6, 0.6, 8, 0.4, f"QUIZ {n}  ·  PAUSE AND SOLVE", size=14, color=AMBER_L, bold=True)
    text(s, 0.6, 1.1, 11.5, 1.0, prompt, size=40, font=HEAD, color=WHITE, bold=True)
    rect(s, 0.6, 2.6, 6.2, 0.62 * len(body_lines) + 0.7, fill="1D2D52", rounded=True, radius=0.05)
    yy = 2.95
    for ln in body_lines:
        text(s, 1.0, yy, 5.6, 0.6, ln, size=30, font=HEAD, color=WHITE, anchor="m")
        yy += 0.62
    if hint:
        text(s, 7.3, 2.7, 5.4, 2.5, hint, size=18, color=LIGHT, spacing=1.25)
    return s


def check(slide, x, y, lines, size=15):
    text(slide, x, y, 5.5, 0.3, "CHECK", size=11, color=TEAL, bold=True)
    text(slide, x, y + 0.3, 5.5, 1.2, lines, size=size, color=GREY, font=HEAD, after=3)


def mini_axes(slide, x, y, size=1.7, kind="point"):
    """Small coordinate picture for the solution set of a 2-variable homogeneous system."""
    cx, cy = x + size / 2, y + size / 2
    if kind == "plane":
        rect(slide, x, y, size, size, fill="C9CED8")
    line(slide, x, cy, x + size, cy, color=TEAL, lw=1.5)
    line(slide, cx, y, cx, y + size, color=TEAL, lw=1.5)
    text(slide, x + size - 0.25, cy + 0.02, 0.3, 0.3, "a", size=12, color=TEAL, font=HEAD, italic=True)
    text(slide, cx + 0.05, y - 0.02, 0.3, 0.3, "b", size=12, color=TEAL, font=HEAD, italic=True)
    if kind == "point":
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.09), Inches(cy - 0.09), Inches(0.18), Inches(0.18))
        d.shadow.inherit = False; d.fill.solid(); d.fill.fore_color.rgb = rgb(NAVY); d.line.fill.background()
    elif kind == "line":
        line(slide, x + 0.1, y + 0.1, x + size - 0.1, y + size - 0.1, color=NAVY, lw=2.5)


# ============================================================== SLIDES ======
# 1 · Cover ------------------------------------------------------------------
s = new_slide(dark=True)
text(s, 0.6, 0.7, 8, 0.4, "MATH FOR DATA SCIENCE", size=16, color=AMBER_L, bold=True)
text(s, 0.6, 1.1, 8, 0.4, "Linear Algebra  ·  Week 2", size=18, color=LIGHT)
text(s, 0.6, 2.4, 11.0, 2.2, "Solving Systems of Linear Equations", size=66, font=HEAD, color=WHITE, bold=True, spacing=1.0)
text(s, 0.6, 5.0, 10.5, 1.3,
     "Elimination  ·  Matrix row reduction  ·  Row operations that preserve singularity  ·  "
     "Rank of a matrix  ·  Row echelon form  ·  Reduced row echelon form  ·  Gaussian elimination",
     size=18, color=LIGHT, spacing=1.3)

# 2 · Roadmap ----------------------------------------------------------------
s = new_slide("Roadmap for this week", "Week 2")
topics = [
    ("Solving systems of equations", "Non-singular, singular, and systems with more variables."),
    ("Matrix row reduction", "Turning a system into a matrix and reducing its rows."),
    ("Row operations that preserve singularity", "Swap, scale, add, and what happens to the determinant."),
    ("Rank of a matrix", "How much information a system really carries."),
    ("Row echelon form and reduced row echelon form", "Pivots, the staircase, and counting the rank."),
    ("Gaussian elimination", "The complete algorithm on the augmented matrix."),
]
for i, (t, b) in enumerate(topics):
    col, row = i % 2, i // 2
    card(s, 0.6 + col * 6.2, 1.75 + row * 1.55, 5.95, 1.35, t, b, number=i + 1)
text(s, 0.6, 6.45, 12, 0.4, "Three short in-class quizzes are built into the session: after non-singular systems, after singular systems, and after rank.",
     size=15, color=GREY, italic=True)

# ===================================================== PART 1 · SYSTEMS =====
divider(1, "Solving systems of linear equations",
        "From two equations in two unknowns to three equations in three unknowns, using only three legal moves.")

# 4 · Definition: linear system ----------------------------------------------
s = new_slide("Linear systems", "Solving systems of equations")
defbox(s, 0.6, 1.7, 7.2, 4.6, [
    [("A linear equation", {"bold": True}), (" in the variables a, b, c, … is an equation of the form", {})],
    [("α₁a + α₂b + α₃c + … = β,", {"font": HEAD, "size": 20})],
    "where the coefficients α₁, α₂, … and the constant β are given numbers. No products or powers of the variables appear.",
    [("A system of linear equations", {"bold": True}), (" is a finite collection of linear equations in the same variables, considered together.", {})],
    [("A solution", {"bold": True}), (" of the system is an assignment of numbers to the variables that satisfies every equation at the same time.", {})],
])
text(s, 8.3, 1.7, 4.4, 0.35, "EXAMPLES", size=12, color=TEAL, bold=True)
eqs(s, 8.3, 2.1, ["a + b = 10", "a + 2b = 12"], w=4.4)
text(s, 8.3, 3.1, 4.4, 0.6, "two equations, two unknowns; the solution is a = 8, b = 2", size=14, color=GREY)
eqs(s, 8.3, 3.9, ["a + b + 2c = 12", "3a − 3b − c = 3", "2a − b + 6c = 24"], w=4.4)
text(s, 8.3, 5.35, 4.4, 0.6, "three equations, three unknowns; we solve it later today", size=14, color=GREY)

# 5 · A system and its solution ----------------------------------------------
s = new_slide("A system and its solution", "Solving systems of equations")
text(s, 0.6, 1.65, 12, 0.5, "An apple and a banana cost $10. An apple and two bananas cost $12. What does each cost?",
     size=20, color=GREY)
rect(s, 0.6, 2.5, 3.6, 2.6, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 2.7, 3.2, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 3.1, ["a + b = 10", "a + 2b = 12"], size=26, gap=0.6, w=3.2)
arrow(s, 4.6, 3.55, w=3.6, h=0.4)
text(s, 4.5, 2.6, 3.8, 0.9, "Some process:\nmanipulating the equations", size=17, color=NAVY, align="c", bold=True)
text(s, 4.5, 4.1, 3.8, 1.0, "swapping equations\nadding equations\nmultiplying an equation by a constant", size=14, color=GREY, align="c")
rect(s, 8.6, 2.5, 3.6, 2.6, fill=RES, rounded=True, radius=0.05)
text(s, 8.9, 2.7, 3.2, 0.3, "SOLVED SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 8.9, 3.1, ["a = 8", "b = 2"], size=26, gap=0.6, w=3.2)
text(s, 0.6, 5.5, 12, 0.9,
     "Reading it off: the second purchase differs from the first by exactly one banana, so a banana costs 12 − 10 = 2, "
     "and the apple costs 10 − 2 = 8. Elimination turns this intuition into a procedure.",
     size=16, color=GREY, spacing=1.2)

# 6 · Equivalent systems and legal operations --------------------------------
s = new_slide("Equivalent systems and the three legal operations", "Solving systems of equations")
defbox(s, 0.6, 1.65, 12.1, 2.5, [
    [("Two systems are ", {}), ("equivalent", {"bold": True}), (" if they have exactly the same solutions. "
     "Each of the following operations replaces a system by an equivalent one:", {})],
    [("(1) ", {"bold": True}), ("swap two equations;   ", {}),
     ("(2) ", {"bold": True}), ("multiply an equation by a non-zero constant;   ", {}),
     ("(3) ", {"bold": True}), ("add a multiple of one equation to another.", {})],
    "Solving a system means applying these operations until each equation involves a single variable.",
])
rect(s, 0.6, 4.45, 5.9, 2.2, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 4.6, 5.3, 0.3, "MULTIPLYING BY A CONSTANT", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 4.95, ["a + b = 10", "× 7", [("7a + 7b = 70", {"bold": True})]], size=22, gap=0.5, w=5.3)
rect(s, 6.8, 4.45, 5.9, 2.2, fill=TINT, rounded=True, radius=0.05)
text(s, 7.1, 4.6, 5.3, 0.3, "ADDING TWO EQUATIONS", size=12, color=TEAL, bold=True)
eqs(s, 7.1, 4.95, ["a + b = 10", "+ (2a + 3b = 22)", [("3a + 4b = 32", {"bold": True})]], size=22, gap=0.5, w=5.3)

# 7 · Worked example 2×2 ------------------------------------------------------
s = new_slide("Worked example: two equations, two unknowns", "Solving systems of equations")
steps = [
    ("System", ["5a + b = 17", "4a − 3b = 6"], ""),
    ("Divide by the coefficient of a", ["a + 0.2b = 3.4", "a − 0.75b = 1.5"], "Both equations now start with a."),
    ("Subtract equation 1 from equation 2", ["−0.95b = −1.9", [("b = 2", {"bold": True})]], "a is eliminated from equation 2."),
    ("Back-substitute b = 2", ["a + 0.2(2) = 3.4", [("a = 3", {"bold": True})]], "Solution: a = 3, b = 2."),
]
for i, (t, lines, note) in enumerate(steps):
    x = 0.6 + i * 3.1
    rect(s, x, 1.75, 2.9, 3.6, fill=RES if i == 3 else TINT, rounded=True, radius=0.05)
    card_num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.22), Inches(1.95), Inches(0.4), Inches(0.4))
    card_num.shadow.inherit = False; card_num.fill.solid(); card_num.fill.fore_color.rgb = rgb(TEAL); card_num.line.fill.background()
    tf = card_num.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(i + 1); r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = rgb(WHITE); r.font.name = BODY
    text(s, x + 0.72, 1.9, 2.1, 0.8, t, size=14, bold=True, anchor="m")
    eqs(s, x + 0.25, 2.9, lines, size=22, gap=0.55, w=2.5)
    text(s, x + 0.25, 4.3, 2.5, 0.9, note, size=13, color=GREY)
text(s, 0.6, 5.6, 12, 1.0,
     "The pattern: make the leading coefficients equal, subtract to eliminate a variable, solve the smaller system, then substitute back. "
     "Every step is one of the three legal operations, so the final system is equivalent to the original.",
     size=16, color=GREY, spacing=1.2)

# 8 · Zero coefficient --------------------------------------------------------
s = new_slide("What if a coefficient of a is already zero?", "Solving systems of equations")
rect(s, 0.6, 1.75, 3.7, 2.4, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 1.9, 3.1, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 2.3, ["5a + b = 17", "3b = 6"], size=26, gap=0.62, w=3.1)
arrow(s, 4.6, 2.75, w=1.0, h=0.4)
rect(s, 5.9, 1.75, 3.2, 2.4, fill=TINT, rounded=True, radius=0.05)
text(s, 6.2, 1.9, 2.8, 0.3, "ISOLATE b", size=12, color=TEAL, bold=True)
eqs(s, 6.2, 2.3, ["3b = 6", [("b = 2", {"bold": True})]], size=26, gap=0.62, w=2.8)
arrow(s, 9.4, 2.75, w=1.0, h=0.4)
rect(s, 10.7, 1.75, 2.0, 2.4, fill=RES, rounded=True, radius=0.05)
text(s, 10.95, 1.9, 1.7, 0.3, "SOLVED", size=12, color=TEAL, bold=True)
eqs(s, 10.95, 2.3, ["a = 3", "b = 2"], size=26, gap=0.62, w=1.7)
text(s, 0.6, 4.5, 12, 1.8, [
    "There is nothing to eliminate: a already has coefficient 0 in the second equation, so that equation gives b directly.",
    "Then a + 0.2(2) = 3.4 gives a + 0.4 = 3.4, so a = 3.",
    "In matrix language this zero is a head start: the row is already in the shape elimination is trying to produce.",
], size=17, color=GREY, spacing=1.2, after=8)

# 9 · Quiz 1 -----------------------------------------------------------------
quiz(1, "Solve the following system of equations", ["2a + 5b = 46", "8a + b = 32"],
     hint="Use the three legal operations: scale an equation, subtract to eliminate one variable, then back-substitute.\n\n"
          "Write your answer as a pair (a, b).")

# 10 · Solution 1 -------------------------------------------------------------
s = new_slide("Quiz 1 · Solution", "Solving systems of equations")
rect(s, 0.6, 1.7, 3.6, 2.0, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 1.85, 3.1, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 2.2, ["2a + 5b = 46", "8a + b = 32"], size=24, gap=0.58, w=3.1)
text(s, 4.6, 1.7, 8.1, 0.3, "ELIMINATE b", size=12, color=TEAL, bold=True)
eqs(s, 4.6, 2.05, [
    "Multiply equation 2 by 5:          40a + 5b = 160",
    "Subtract equation 1:                  38a = 114",
    [("a = 3", {"bold": True})],
], size=20, gap=0.5, w=8.1)
text(s, 4.6, 3.7, 8.1, 0.3, "BACK-SUBSTITUTE", size=12, color=TEAL, bold=True)
eqs(s, 4.6, 4.05, ["8(3) + b = 32", [("b = 8", {"bold": True})]], size=20, gap=0.5, w=8.1)
rect(s, 0.6, 4.0, 3.6, 1.3, fill=RES, rounded=True, radius=0.05)
text(s, 0.9, 4.15, 3.1, 0.3, "SOLUTION", size=12, color=TEAL, bold=True)
text(s, 0.9, 4.45, 3.1, 0.7, "a = 3,  b = 8", size=28, font=HEAD, bold=True)
check(s, 4.6, 5.3, ["2(3) + 5(8) = 6 + 40 = 46  ✓", "8(3) + 8 = 24 + 8 = 32  ✓"])

# 11 · Singular, redundant ----------------------------------------------------
s = new_slide("What if the system is singular? (redundant)", "Solving systems of equations")
rect(s, 0.6, 1.75, 3.5, 2.3, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 1.9, 3.0, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 2.3, ["a + b = 10", "2a + 2b = 20"], size=24, gap=0.6, w=3.0)
arrow(s, 4.35, 2.7, w=0.9, h=0.38, label="divide eq. 2 by 2")
rect(s, 5.5, 1.75, 3.3, 2.3, fill=TINT, rounded=True, radius=0.05)
eqs(s, 5.8, 2.3, ["a + b = 10", "a + b = 10"], size=24, gap=0.6, w=2.8)
arrow(s, 9.05, 2.7, w=0.9, h=0.38, label="subtract eq. 1")
rect(s, 10.2, 1.75, 2.5, 2.3, fill=PIV, rounded=True, radius=0.05)
eqs(s, 10.5, 2.3, ["a + b = 10", [("0 = 0", {"bold": True})]], size=24, gap=0.6, w=2.0)
text(s, 0.6, 4.6, 7.4, 2.0, [
    [("The second equation carried no new information.", {"bold": True})],
    "After elimination only one equation survives, and one equation cannot pin down two unknowns.",
    "Pick any value a = x. Then b = 10 − x also works: the system has infinitely many solutions, with one degree of freedom.",
], size=17, color=GREY, spacing=1.2, after=8)
rect(s, 8.4, 4.6, 4.3, 1.9, fill=RES, rounded=True, radius=0.05)
text(s, 8.7, 4.75, 3.8, 0.3, "SOLUTION SET", size=12, color=TEAL, bold=True)
eqs(s, 8.7, 5.1, ["a = x", "b = 10 − x"], size=24, gap=0.55, w=3.8)
text(s, 8.7, 6.15, 3.8, 0.3, "for every real number x", size=14, color=GREY)

# 12 · Singular, contradictory ------------------------------------------------
s = new_slide("What if the system is singular? (contradictory)", "Solving systems of equations")
rect(s, 0.6, 1.75, 3.5, 2.3, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 1.9, 3.0, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 2.3, ["a + b = 10", "2a + 2b = 24"], size=24, gap=0.6, w=3.0)
arrow(s, 4.35, 2.7, w=0.9, h=0.38, label="divide eq. 2 by 2")
rect(s, 5.5, 1.75, 3.3, 2.3, fill=TINT, rounded=True, radius=0.05)
eqs(s, 5.8, 2.3, ["a + b = 10", "a + b = 12"], size=24, gap=0.6, w=2.8)
arrow(s, 9.05, 2.7, w=0.9, h=0.38, label="subtract eq. 1")
rect(s, 10.2, 1.75, 2.5, 2.3, fill="F6D5D2", rounded=True, radius=0.05)
eqs(s, 10.5, 2.3, ["a + b = 10", [("0 = 2", {"bold": True, "color": RED})]], size=24, gap=0.6, w=2.0)
text(s, 0.6, 4.6, 7.4, 2.0, [
    [("A contradiction.", {"bold": True, "color": RED})],
    "No choice of a and b can make 0 equal to 2, so no pair satisfies both equations.",
    "The left-hand sides are proportional (2a + 2b is twice a + b) but the right-hand sides are not (24 is not twice 10).",
], size=17, color=GREY, spacing=1.2, after=8)
rect(s, 8.4, 4.6, 4.3, 1.9, fill="F6D5D2", rounded=True, radius=0.05)
text(s, 8.7, 4.75, 3.8, 0.3, "SOLUTION SET", size=12, color=RED, bold=True)
text(s, 8.7, 5.1, 3.8, 0.7, "empty", size=28, font=HEAD, bold=True)
text(s, 8.7, 5.85, 3.8, 0.5, "the system has no solutions", size=14, color=GREY)

# 13 · Classifying systems ---------------------------------------------------
s = new_slide("Classifying systems", "Solving systems of equations")
defbox(s, 0.6, 1.65, 12.1, 1.95, [
    [("A system is ", {}), ("consistent", {"bold": True}), (" if it has at least one solution and ", {}),
     ("inconsistent", {"bold": True}), (" if it has none.", {})],
    [("A square system (as many equations as unknowns) is ", {}), ("non-singular", {"bold": True}),
     (" if it has exactly one solution, and ", {}), ("singular", {"bold": True}),
     (" otherwise, that is, when it has infinitely many solutions or none.", {})],
])
cases = [
    ("Unique solution", "non-singular", ["a + b = 10", "a + 2b = 12"], "two lines that cross in one point", RES),
    ("Infinitely many solutions", "singular · redundant", ["a + b = 10", "2a + 2b = 20"], "two equations for the same line", PIV),
    ("No solution", "singular · contradictory", ["a + b = 10", "2a + 2b = 24"], "two parallel lines", "F6D5D2"),
]
for i, (t, sub, lines, geo, fill) in enumerate(cases):
    x = 0.6 + i * 4.1
    rect(s, x, 3.9, 3.9, 2.75, fill=fill, rounded=True, radius=0.05)
    text(s, x + 0.25, 4.05, 3.4, 0.4, t, size=18, bold=True)
    text(s, x + 0.25, 4.42, 3.4, 0.3, sub.upper(), size=11, color=TEAL, bold=True)
    eqs(s, x + 0.25, 4.85, lines, size=20, gap=0.45, w=3.4)
    text(s, x + 0.25, 5.9, 3.4, 0.6, "Geometrically: " + geo, size=13, color=GREY)

# 14 · Quiz 2 -----------------------------------------------------------------
quiz(2, "Solve the following system of equations", ["5a + b = 11", "10a + 2b = 22"],
     hint="Look closely at the two equations before you start computing.\n\n"
          "Is the system non-singular, redundant, or contradictory? Describe every solution.")

# 15 · Solution 2 -------------------------------------------------------------
s = new_slide("Quiz 2 · Solution", "Solving systems of equations")
rect(s, 0.6, 1.7, 3.6, 2.0, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 1.85, 3.1, 0.3, "SYSTEM", size=12, color=TEAL, bold=True)
eqs(s, 0.9, 2.2, ["5a + b = 11", "10a + 2b = 22"], size=24, gap=0.58, w=3.1)
text(s, 4.6, 1.7, 8.1, 0.3, "ELIMINATE", size=12, color=TEAL, bold=True)
eqs(s, 4.6, 2.05, [
    "Divide equation 2 by 2:       5a + b = 11",
    "This is equation 1 again. Subtracting gives",
    [("0 = 0", {"bold": True})],
], size=20, gap=0.5, w=8.1)
text(s, 4.6, 3.8, 8.1, 1.6, [
    [("The system is singular and redundant: ", {"bold": True, "color": NAVY}),
     ("equation 2 is just twice equation 1, so it adds no information.", {})],
    "One equation in two unknowns leaves one degree of freedom, so there are infinitely many solutions.",
], size=17, color=GREY, spacing=1.2, after=8)
rect(s, 0.6, 4.0, 3.6, 1.6, fill=RES, rounded=True, radius=0.05)
text(s, 0.9, 4.15, 3.1, 0.3, "SOLUTION SET", size=12, color=TEAL, bold=True)
text(s, 0.9, 4.45, 3.1, 0.5, "b = 11 − 5a", size=26, font=HEAD, bold=True)
text(s, 0.9, 5.05, 3.1, 0.4, "for any real number a", size=14, color=GREY)
check(s, 4.6, 5.4, ["(a, b) = (0, 11), (1, 6), (2, 1), (−1, 16), … all satisfy both equations."])

# 16 · Elimination with three variables, step 1 ------------------------------
s = new_slide("Elimination with three variables (1 of 3)", "Solving systems of equations")
cols = [
    ("SYSTEM", ["a + b + 2c = 12", "3a − 3b − c = 3", "2a − b + 6c = 24"], "Three equations, three unknowns.", TINT),
    ("DIVIDE EACH ROW BY ITS COEFFICIENT OF a", ["a + b + 2c = 12", "a − b − ⅓c = 1", "a − ½b + 3c = 12"], "Now every equation starts with a.", TINT),
    ("SUBTRACT EQUATION 1 FROM THE OTHERS", ["a + b + 2c = 12", [("−2b − 7⁄3 c = −11", {"bold": True})], [("−3⁄2 b + c = 0", {"bold": True})]],
     "a is isolated in equation 1; equations 2 and 3 form a 2×2 system in b and c.", RES),
]
for i, (t, lines, note, fill) in enumerate(cols):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.9, 3.9, fill=fill, rounded=True, radius=0.05)
    text(s, x + 0.25, 1.9, 3.4, 0.6, t, size=11, color=TEAL, bold=True)
    eqs(s, x + 0.25, 2.55, lines, size=21, gap=0.6, w=3.4)
    text(s, x + 0.25, 4.5, 3.4, 1.0, note, size=14, color=GREY, spacing=1.15)
text(s, 0.6, 5.95, 12, 0.8, "Same idea as before: use the first equation to remove a from every other equation, then solve what is left.",
     size=16, color=GREY, spacing=1.2)

# 17 · step 2 -----------------------------------------------------------------
s = new_slide("Elimination with three variables (2 of 3)", "Solving systems of equations")
cols = [
    ("THE SYSTEM SO FAR", ["a + b + 2c = 12", "−2b − 7⁄3 c = −11", "−3⁄2 b + c = 0"], "Work on the last two equations only.", TINT),
    ("DIVIDE BY THE COEFFICIENT OF b", ["a + b + 2c = 12", "b + 7⁄6 c = 11⁄2", "b − 2⁄3 c = 0"], "Both now start with b.", TINT),
    ("SUBTRACT EQUATION 2 FROM EQUATION 3", ["a + b + 2c = 12", "b + 7⁄6 c = 11⁄2", [("−11⁄6 c = −11⁄2", {"bold": True})]],
     "b is isolated in equation 2, and the last equation gives c = 3.", RES),
]
for i, (t, lines, note, fill) in enumerate(cols):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.9, 3.9, fill=fill, rounded=True, radius=0.05)
    text(s, x + 0.25, 1.9, 3.4, 0.6, t, size=11, color=TEAL, bold=True)
    eqs(s, x + 0.25, 2.55, lines, size=21, gap=0.6, w=3.4)
    text(s, x + 0.25, 4.5, 3.4, 1.0, note, size=14, color=GREY, spacing=1.15)
text(s, 0.6, 5.95, 12, 0.8, "The system is now triangular: the first equation has a, b, c; the second has b, c; the third has c alone.",
     size=16, color=GREY, spacing=1.2)

# 18 · step 3 -----------------------------------------------------------------
s = new_slide("Three variables (3 of 3): back substitution", "Solving systems of equations")
rect(s, 0.6, 1.75, 3.9, 2.6, fill=TINT, rounded=True, radius=0.05)
text(s, 0.85, 1.9, 3.4, 0.3, "TRIANGULAR SYSTEM", size=11, color=TEAL, bold=True)
eqs(s, 0.85, 2.3, ["a + b + 2c = 12", "b + 7⁄6 c = 11⁄2", "c = 3"], size=22, gap=0.6, w=3.4)
text(s, 4.9, 1.75, 7.8, 0.3, "SUBSTITUTE UPWARDS", size=11, color=TEAL, bold=True)
eqs(s, 4.9, 2.1, [
    "c = 3 into equation 2:   b + 7⁄6 · 3 = 11⁄2,  so  b + 7⁄2 = 11⁄2,  so  b = 2",
    "b = 2, c = 3 into equation 1:   a + 2 + 6 = 12,  so  a = 4",
], size=18, gap=0.7, w=7.8)
rect(s, 4.9, 3.75, 7.8, 1.0, fill=RES, rounded=True, radius=0.05)
text(s, 5.2, 3.85, 7.2, 0.8, "Solution:  a = 4,  b = 2,  c = 3", size=26, font=HEAD, bold=True, anchor="m")
check(s, 4.9, 5.0, ["4 + 2 + 2(3) = 12  ✓", "3(4) − 3(2) − 3 = 3  ✓", "2(4) − 2 + 6(3) = 24  ✓"])
text(s, 0.6, 4.7, 3.9, 1.8, "Back substitution: solve the last equation first and feed each answer into the equation above it.",
     size=15, color=GREY, spacing=1.2)

# ================================================ PART 2 · ROW REDUCTION ====
divider(2, "Matrix row reduction",
        "The same elimination steps, written on the coefficients only. Rows replace equations.")

# 20 · From systems to matrices ----------------------------------------------
s = new_slide("From systems to matrices", "Matrix row reduction")
defbox(s, 0.6, 1.6, 12.1, 1.5, [
    [("The coefficient matrix", {"bold": True}),
     (" of a system lists the coefficients of the variables: one row per equation, one column per variable. ", {}),
     ("Row reduction", {"bold": True}), (" applies the three legal operations to the rows of the matrix instead of to the equations.", {})],
])
# equations row
eqs(s, 0.9, 3.35, ["5a + b = 17", "4a − 3b = 6"], size=20, gap=0.45, w=3.2, title="ORIGINAL SYSTEM")
eqs(s, 5.0, 3.35, ["a + 0.2b = 3.4", "b = 2"], size=20, gap=0.45, w=3.2, title="INTERMEDIATE SYSTEM")
eqs(s, 9.3, 3.35, ["1a + 0b = 3", "0a + 1b = 2"], size=20, gap=0.45, w=3.2, title="SOLVED SYSTEM")
# matrices row
matrix(s, 1.2, 5.0, [["5", "1"], ["4", "−3"]], label="Original matrix")
arrow(s, 3.1, 5.35, w=1.2, h=0.36)
matrix(s, 5.3, 5.0, [["1", "0.2"], ["0", "1"]], label="Upper triangular: row echelon form")
arrow(s, 7.3, 5.35, w=1.2, h=0.36)
matrix(s, 9.6, 5.0, [["1", "0"], ["0", "1"]], label="Diagonal: reduced row echelon form", fill=RES)

# 21 · Singular systems as matrices ------------------------------------------
s = new_slide("Singular systems as matrices", "Matrix row reduction")
rowsdata = [
    (["a + b = 10", "2a + 2b = 20"], [["1", "1"], ["2", "2"]], ["a + b = 10", "0a + 0b = 0"], [["1", "1"], ["0", "0"]], "redundant"),
    (["5a + b = 11", "10a + 2b = 22"], [["5", "1"], ["10", "2"]], ["a + 0.2b = 2.2", "0a + 0b = 0"], [["1", "0.2"], ["0", "0"]], "redundant"),
    (["0a + 0b = 0", "0a + 0b = 0"], [["0", "0"], ["0", "0"]], ["0a + 0b = 0", "0a + 0b = 0"], [["0", "0"], ["0", "0"]], "no information at all"),
]
text(s, 0.6, 1.55, 3, 0.3, "ORIGINAL SYSTEM", size=11, color=TEAL, bold=True)
text(s, 3.6, 1.55, 3, 0.3, "ORIGINAL MATRIX", size=11, color=TEAL, bold=True)
text(s, 7.2, 1.55, 3, 0.3, "AFTER ELIMINATION", size=11, color=TEAL, bold=True)
text(s, 10.3, 1.55, 3, 0.3, "ROW ECHELON FORM", size=11, color=TEAL, bold=True)
for i, (e1, m1, e2, m2, tag) in enumerate(rowsdata):
    y = 1.95 + i * 1.55
    eqs(s, 0.6, y, e1, size=18, gap=0.42, w=3.0)
    matrix(s, 3.7, y + 0.05, m1, cw=0.5, ch=0.42, size=15)
    arrow(s, 5.9, y + 0.3, w=1.0, h=0.34)
    eqs(s, 7.2, y, e2, size=18, gap=0.42, w=3.0)
    matrix(s, 10.4, y + 0.05, m2, cw=0.5, ch=0.42, size=15, hl={(1, 0): PIV, (1, 1): PIV})
    text(s, 12.0, y + 0.2, 1.2, 0.5, tag, size=11, color=GREY, italic=True)
text(s, 0.6, 6.5, 12, 0.45, "A row of zeros appears in the row echelon form exactly when the system is singular.",
     size=17, color=NAVY, bold=True)

# 22 · Elementary row operations ---------------------------------------------
s = new_slide("Elementary row operations", "Row operations that preserve singularity")
defbox(s, 0.6, 1.6, 12.1, 2.05, [
    [("The three ", {}), ("elementary row operations", {"bold": True}), (" on a matrix are: swapping two rows, "
     "multiplying a row by a non-zero scalar, and adding a multiple of one row to another row.", {})],
    [("Theorem. ", {"bold": True}), ("Elementary row operations do not change the solution set of the corresponding system, "
     "and they do not change whether the matrix is singular.", {})],
])
ops = [
    ("Swap", "Rᵢ ↔ Rⱼ", "Reorders the equations. Nothing about the solutions changes."),
    ("Scale", "Rᵢ ← k · Rᵢ,  k ≠ 0", "Multiplies one equation by a constant. Dividing by k undoes it."),
    ("Replace", "Rᵢ ← Rᵢ + k · Rⱼ", "Adds a multiple of another equation. Subtracting k · Rⱼ undoes it."),
]
for i, (t, formula, note) in enumerate(ops):
    x = 0.6 + i * 4.1
    rect(s, x, 3.9, 3.9, 2.75, fill=TINT, rounded=True, radius=0.05)
    text(s, x + 0.25, 4.05, 3.4, 0.4, t, size=18, bold=True)
    text(s, x + 0.25, 4.5, 3.4, 0.6, formula, size=24, font=HEAD, color=TEAL)
    text(s, x + 0.25, 5.25, 3.4, 1.0, note, size=14, color=GREY, spacing=1.15)
    text(s, x + 0.25, 6.2, 3.4, 0.4, "reversible", size=12, color=TEAL, bold=True)

# 23 · Row operations and the determinant ------------------------------------
s = new_slide("Row operations and the determinant", "Row operations that preserve singularity")
text(s, 0.6, 1.55, 12, 0.5, "Start from the same matrix each time and watch what happens to det = 5 · 3 − 1 · 4 = 11.",
     size=17, color=GREY)
dets = [
    ("Switch the rows", [["4", "3"], ["5", "1"]], "det = 4 · 1 − 3 · 5 = −11", "The sign flips."),
    ("Multiply row 1 by 10", [["50", "10"], ["4", "3"]], "det = 50 · 3 − 10 · 4 = 110 = 10 · 11", "The determinant is multiplied by 10."),
    ("Add row 1 to row 2", [["5", "1"], ["9", "4"]], "det = 5 · 4 − 1 · 9 = 11", "The determinant does not change."),
]
for i, (t, m, d, note) in enumerate(dets):
    x = 0.6 + i * 4.1
    rect(s, x, 2.15, 3.9, 4.0, fill=TINT, rounded=True, radius=0.05)
    text(s, x + 0.25, 2.3, 3.4, 0.4, t, size=16, bold=True)
    matrix(s, x + 0.35, 2.85, [["5", "1"], ["4", "3"]], cw=0.48, ch=0.42, size=15)
    arrow(s, x + 1.75, 3.1, w=0.5, h=0.3)
    matrix(s, x + 2.4, 2.85, m, cw=0.48, ch=0.42, size=15, fill=RES)
    text(s, x + 0.25, 4.1, 3.4, 0.5, "det = 11", size=13, color=GREY)
    text(s, x + 0.25, 4.6, 3.4, 0.5, d, size=16, font=HEAD, color=NAVY)
    text(s, x + 0.25, 5.3, 3.4, 0.7, note, size=14, color=GREY)
text(s, 0.6, 6.4, 12, 0.5, "A determinant that is zero stays zero, and one that is non-zero stays non-zero: row operations preserve singularity.",
     size=16, color=NAVY, bold=True)

# ========================================================= PART 3 · RANK ====
divider(3, "Rank of a matrix",
        "Two equations do not always carry two pieces of information. Rank measures how much they really say.")

# 25 · Systems of information ------------------------------------------------
s = new_slide("Systems of information", "Rank of a matrix")
info = [
    ("System 1", ["The dog is black.", "The cat is orange."], "Two sentences", "Two pieces of information", "Rank = 2", RES),
    ("System 2", ["The dog is black.", "The dog is black."], "Two sentences", "One piece of information", "Rank = 1", PIV),
    ("System 3", ["The dog.", "The dog."], "Two sentences", "Zero pieces of information", "Rank = 0", "F6D5D2"),
]
for i, (t, sent, a, b, r, fill) in enumerate(info):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.9, 4.6, fill=fill, rounded=True, radius=0.05)
    text(s, x + 0.25, 1.9, 3.4, 0.4, t, size=18, bold=True)
    text(s, x + 0.25, 2.45, 3.4, 1.2, sent, size=20, font=HEAD, color=NAVY, after=6)
    text(s, x + 0.25, 3.9, 3.4, 0.4, a, size=15, color=GREY)
    text(s, x + 0.25, 4.3, 3.4, 0.5, b, size=16, color=TEAL, bold=True)
    text(s, x + 0.25, 5.3, 3.4, 0.7, r, size=30, font=HEAD, bold=True)
text(s, 0.6, 6.5, 12, 0.45, "Rank counts independent pieces of information, not the number of statements.", size=16, color=GREY, italic=True)

# 26 · Rank of a 2×2 system --------------------------------------------------
s = new_slide("Rank of a 2×2 system", "Rank of a matrix")
sys2 = [
    ("System 1", ["a + b = 0", "a + 2b = 0"], [["1", "1"], ["1", "2"]], "Two pieces of information", "Rank = 2", RES),
    ("System 2", ["a + b = 0", "2a + 2b = 0"], [["1", "1"], ["2", "2"]], "One piece of information", "Rank = 1", PIV),
    ("System 3", ["0a + 0b = 0", "0a + 0b = 0"], [["0", "0"], ["0", "0"]], "Zero pieces of information", "Rank = 0", "F6D5D2"),
]
for i, (t, lines, m, b, r, fill) in enumerate(sys2):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.9, 4.6, fill=fill, rounded=True, radius=0.05)
    text(s, x + 0.25, 1.9, 3.4, 0.4, t, size=18, bold=True)
    eqs(s, x + 0.25, 2.4, lines, size=20, gap=0.45, w=2.0)
    matrix(s, x + 2.3, 2.45, m, cw=0.5, ch=0.42, size=15)
    text(s, x + 0.25, 3.75, 3.4, 0.4, "Two equations", size=15, color=GREY)
    text(s, x + 0.25, 4.15, 3.4, 0.5, b, size=16, color=TEAL, bold=True)
    text(s, x + 0.25, 5.3, 3.4, 0.7, r, size=30, font=HEAD, bold=True)
text(s, 0.6, 6.5, 12, 0.45, "The rank of the matrix is the rank of the system it represents.", size=16, color=GREY, italic=True)

# 27 · Rank and the solution space -------------------------------------------
s = new_slide("Rank and the solution space", "Rank of a matrix")
defbox(s, 0.6, 1.55, 12.1, 1.55, [
    [("For a homogeneous system (all constants equal to 0) the ", {}), ("solution space", {"bold": True}),
     (" is the set of all its solutions; its ", {}), ("dimension", {"bold": True}),
     (" is the number of free parameters needed to describe it. ", {}),
     ("Rank of the matrix = number of variables − dimension of the solution space.", {"bold": True})],
])
geo = [
    ([["1", "1"], ["1", "2"]], "Rank = 2", "dimension 0", "only the point (0, 0)", "point"),
    ([["1", "1"], ["2", "2"]], "Rank = 1", "dimension 1", "the line b = −a", "line"),
    ([["0", "0"], ["0", "0"]], "Rank = 0", "dimension 2", "the whole plane", "plane"),
]
for i, (m, r, d, desc, kind) in enumerate(geo):
    x = 0.6 + i * 4.1
    matrix(s, x + 0.2, 3.4, m, cw=0.48, ch=0.42, size=15)
    text(s, x + 1.7, 3.4, 2.2, 0.5, r, size=22, font=HEAD, bold=True)
    text(s, x + 1.7, 3.9, 2.2, 0.4, "solution space has " + d, size=13, color=GREY)
    mini_axes(s, x + 0.35, 4.55, size=1.5, kind=kind)
    text(s, x + 2.1, 4.9, 1.8, 0.9, desc, size=15, color=NAVY)
text(s, 0.6, 6.4, 12, 0.5, "With two variables:  Rank = 2 − (dimension of the solution space).", size=18, font=HEAD, color=NAVY, bold=True)

# 28 · Rank and singularity --------------------------------------------------
s = new_slide("Rank and singularity", "Rank of a matrix")
rs = [([["1", "1"], ["1", "2"]], "Rank = 2", "Non-singular", RES), ([["1", "1"], ["2", "2"]], "Rank = 1", "Singular", PIV),
      ([["0", "0"], ["0", "0"]], "Rank = 0", "Singular", "F6D5D2")]
for i, (m, r, tag, fill) in enumerate(rs):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.9, 3.3, fill=fill, rounded=True, radius=0.05)
    matrix(s, x + 1.25, 2.05, m, cw=0.6, ch=0.52, size=18)
    text(s, x + 0.25, 3.35, 3.4, 0.6, r, size=26, font=HEAD, bold=True, align="c")
    text(s, x + 0.25, 4.1, 3.4, 0.6, tag, size=20, color=TEAL, bold=True, align="c")
defbox(s, 0.6, 5.35, 12.1, 1.3, [
    [("A square n × n matrix is non-singular exactly when its rank equals n (", {}), ("full rank", {"bold": True}),
     ("). Any smaller rank means at least one row carries no new information: the matrix is singular.", {})],
], label="Fact")

# 29 · Quiz 3 -----------------------------------------------------------------
s = new_slide(dark=True)
text(s, 0.6, 0.6, 8, 0.4, "QUIZ 3  ·  PAUSE AND SOLVE", size=14, color=AMBER_L, bold=True)
text(s, 0.6, 1.1, 11.5, 1.0, "Find the rank of each matrix", size=40, font=HEAD, color=WHITE, bold=True)
text(s, 0.6, 2.6, 3, 0.4, "MATRIX 1", size=14, color=LIGHT, bold=True)
matrix(s, 0.6, 3.1, [["5", "1"], ["−1", "3"]], cw=0.75, ch=0.62, size=24, fill="1D2D52", tcolor=WHITE, bcolor=WHITE)
text(s, 4.6, 2.6, 3, 0.4, "MATRIX 2", size=14, color=LIGHT, bold=True)
matrix(s, 4.6, 3.1, [["2", "−1"], ["−6", "3"]], cw=0.75, ch=0.62, size=24, fill="1D2D52", tcolor=WHITE, bcolor=WHITE)
text(s, 8.3, 2.7, 4.4, 2.5, "Hint: write the homogeneous system (constants 0), find its solution space, and use\n\nRank = 2 − dimension of the solution space.",
     size=18, color=LIGHT, spacing=1.25)
# 30 · Solution 3 -------------------------------------------------------------
s = new_slide("Quiz 3 · Solution", "Rank of a matrix")
rect(s, 0.6, 1.7, 5.9, 4.8, fill=RES, rounded=True, radius=0.05)
text(s, 0.9, 1.85, 5.3, 0.4, "Matrix 1", size=18, bold=True)
matrix(s, 0.9, 2.4, [["5", "1"], ["−1", "3"]], cw=0.5, ch=0.42, size=15)
eqs(s, 2.6, 2.4, ["5a + b = 0", "−a + 3b = 0"], size=18, gap=0.42, w=3.6)
text(s, 0.9, 3.5, 5.3, 2.2, [
    "From the first equation b = −5a. Substituting: −a + 3(−5a) = −16a = 0, so a = 0 and then b = 0.",
    "The solution space is the single point (0, 0): dimension 0.",
    [("Rank = 2 − 0 = 2.  Non-singular (det = 15 + 1 = 16 ≠ 0).", {"bold": True, "color": NAVY})],
], size=15, color=GREY, spacing=1.15, after=6)
rect(s, 6.8, 1.7, 5.9, 4.8, fill=PIV, rounded=True, radius=0.05)
text(s, 7.1, 1.85, 5.3, 0.4, "Matrix 2", size=18, bold=True)
matrix(s, 7.1, 2.4, [["2", "−1"], ["−6", "3"]], cw=0.5, ch=0.42, size=15)
eqs(s, 8.8, 2.4, ["2a − b = 0", "−6a + 3b = 0"], size=18, gap=0.42, w=3.6)
text(s, 7.1, 3.5, 5.3, 2.2, [
    "The second row is −3 times the first, so the second equation repeats the first.",
    "Every point on the line b = 2a is a solution: dimension 1.",
    [("Rank = 2 − 1 = 1.  Singular (det = 6 − 6 = 0).", {"bold": True, "color": NAVY})],
], size=15, color=GREY, spacing=1.15, after=6)

# 31 · Rank of a 3×3 system --------------------------------------------------
s = new_slide("Rank of a 3×3 system", "Rank of a matrix")
s3 = [
    ("System 1", ["a + b + c = 0", "a + 2b + c = 0", "a + b + 2c = 0"], [["1", "1", "1"], ["1", "2", "1"], ["1", "1", "2"]], "3 pieces of information", "Rank 3"),
    ("System 2", ["a + b + c = 0", "a + b + 2c = 0", "a + b + 3c = 0"], [["1", "1", "1"], ["1", "1", "2"], ["1", "1", "3"]], "2 pieces of information", "Rank 2"),
    ("System 3", ["a + b + c = 0", "2a + 2b + 2c = 0", "3a + 3b + 3c = 0"], [["1", "1", "1"], ["2", "2", "2"], ["3", "3", "3"]], "1 piece of information", "Rank 1"),
    ("System 4", ["0a + 0b + 0c = 0", "0a + 0b + 0c = 0", "0a + 0b + 0c = 0"], [["0", "0", "0"], ["0", "0", "0"], ["0", "0", "0"]], "0 pieces of information", "Rank 0"),
]
for i, (t, lines, m, b, r) in enumerate(s3):
    x = 0.6 + i * 3.08
    rect(s, x, 1.7, 2.9, 4.95, fill=TINT, rounded=True, radius=0.05)
    text(s, x + 0.2, 1.82, 2.6, 0.4, t, size=16, bold=True)
    eqs(s, x + 0.2, 2.25, lines, size=15, gap=0.36, w=2.6)
    text(s, x + 0.2, 3.4, 2.6, 0.3, "3 equations", size=13, color=GREY)
    text(s, x + 0.2, 3.7, 2.6, 0.4, b, size=14, color=TEAL, bold=True)
    matrix(s, x + 0.5, 4.25, m, cw=0.42, ch=0.36, size=13, gap=0.04)
    text(s, x + 0.2, 5.6, 2.6, 0.6, r, size=24, font=HEAD, bold=True, align="c")

# 32 · Easier way -------------------------------------------------------------
s = new_slide(dark=True)
text(s, 0.6, 1.3, 11, 0.4, "QUESTION", size=16, color=AMBER_L, bold=True)
text(s, 0.6, 1.8, 11.5, 1.0, "An easier way to compute the rank?", size=44, font=HEAD, color=WHITE, bold=True)
text(s, 0.6, 3.4, 11, 0.4, "ANSWER", size=16, color=AMBER_L, bold=True)
text(s, 0.6, 3.9, 11.5, 2.4,
     "Yes. Bring the matrix to row echelon form and count the pivots. Equivalently, count the ones on the diagonal "
     "of its reduced row echelon form. The next part makes both forms precise.",
     size=26, color=LIGHT, spacing=1.25)

# ============================================= PART 4 · ROW ECHELON FORM ====
divider(4, "Row echelon form",
        "A staircase shape that elimination always reaches, and the fastest way to read off the rank.")

# 34 · Definition REF ---------------------------------------------------------
s = new_slide("Row echelon form: definition", "Row echelon form")
defbox(s, 0.6, 1.55, 7.0, 3.3, [
    [("A matrix is in ", {}), ("row echelon form", {"bold": True}), (" if", {})],
    "(1)  all rows consisting entirely of zeros are at the bottom;",
    [("(2)  in each non-zero row the first non-zero entry, called the ", {}), ("pivot", {"bold": True}),
     (", lies strictly to the right of the pivot of the row above.", {})],
    [("The ", {}), ("rank", {"bold": True}), (" of a matrix is the number of pivots in its row echelon form.", {})],
])
text(s, 0.6, 5.05, 7.0, 1.5, [
    [("Note. ", {"bold": True}), ("In general pivots other than 1 are allowed. In this class we divide each row by its pivot so that "
     "every pivot is 1; this makes no mathematical difference.", {})],
], size=15, color=GREY, spacing=1.15)
star = [["2", "*", "*", "*", "*"], ["0", "1", "*", "*", "*"], ["0", "0", "3", "*", "*"], ["0", "0", "0", "−5", "*"], ["0", "0", "0", "0", "1"]]
hl5 = {(i, i): PIV for i in range(5)}
matrix(s, 7.9, 1.7, star, cw=0.4, ch=0.36, size=13, gap=0.04, hl=hl5, label="Rank 5", lsize=15, lcolor=NAVY)
star3 = [["3", "*", "*", "*", "*"], ["0", "0", "1", "*", "*"], ["0", "0", "0", "−4", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
matrix(s, 10.5, 1.7, star3, cw=0.4, ch=0.36, size=13, gap=0.04, hl={(0, 0): PIV, (1, 2): PIV, (2, 3): PIV}, label="Rank 3", lsize=15, lcolor=NAVY)
text(s, 7.9, 4.5, 4.9, 2.0, [
    "•  Zero rows sit at the bottom.",
    "•  Each non-zero row has a pivot (its leftmost non-zero entry).",
    "•  Every pivot is to the right of the pivots above it: a staircase.",
    "•  The rank is the number of pivots: 5 and 3.",
], size=14, color=GREY, spacing=1.15, after=4)

# 35 · REF steps 2×2 ----------------------------------------------------------
s = new_slide("Reaching row echelon form, step by step", "Row echelon form")
matrix(s, 0.8, 2.4, [["5", "1"], ["4", "−3"]], cw=0.62, ch=0.52, size=18, label="Original matrix")
arrow(s, 2.75, 2.75, w=1.1, h=0.36, label="step 1", lsize=14, lw_=1.1)
matrix(s, 4.25, 2.4, [["1", "0.2"], ["1", "−0.75"]], cw=0.72, ch=0.52, size=18)
arrow(s, 6.45, 2.75, w=1.1, h=0.36, label="step 2", lsize=14, lw_=1.1)
matrix(s, 7.95, 2.4, [["1", "0.2"], ["0", "−0.95"]], cw=0.72, ch=0.52, size=18, hl={(1, 0): RES, (1, 1): RES})
arrow(s, 10.15, 2.75, w=1.1, h=0.36, label="step 3", lsize=14, lw_=1.1)
matrix(s, 11.55, 2.4, [["1", "0.2"], ["0", "1"]], cw=0.6, ch=0.52, size=18, fill=RES, label="Row echelon form")
rect(s, 0.6, 4.4, 12.1, 2.25, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 4.55, 11.5, 0.3, "THE STEPS", size=12, color=TEAL, bold=True)
text(s, 0.9, 4.9, 11.5, 1.7, [
    [("Step 1  ", {"bold": True}), ("divide each row by its leftmost coefficient:  5a + b = 17 becomes a + 0.2b = 3.4,  4a − 3b = 6 becomes a − 0.75b = 1.5", {})],
    [("Step 2  ", {"bold": True}), ("subtract row 1 from row 2:  [ 1   −0.75 ]  −  [ 1   0.2 ]  =  [ 0   −0.95 ]", {})],
    [("Step 3  ", {"bold": True}), ("divide row 2 by its leftmost non-zero coefficient:  [ 0   −0.95 ] ÷ (−0.95)  =  [ 0   1 ]", {})],
], size=17, font=HEAD, color=NAVY, after=8)

# 36 · REF, singularity and rank ----------------------------------------------
s = new_slide("Row echelon form, singularity and rank", "Row echelon form")
rr = [("Non-singular matrix", [["5", "1"], ["4", "−3"]], [["1", "0.2"], ["0", "1"]], {(0, 0): PIV, (1, 1): PIV}, "2 pivots", "Rank 2"),
      ("Singular matrix", [["5", "1"], ["10", "2"]], [["1", "0.2"], ["0", "0"]], {(0, 0): PIV}, "1 pivot", "Rank 1"),
      ("Singular matrix", [["0", "0"], ["0", "0"]], [["0", "0"], ["0", "0"]], {}, "0 pivots", "Rank 0")]
for i, (t, m1, m2, hl, p, r) in enumerate(rr):
    y = 1.75 + i * 1.6
    text(s, 0.6, y + 0.3, 3.2, 0.5, t, size=18, bold=True)
    matrix(s, 3.9, y, m1, cw=0.56, ch=0.48, size=16)
    arrow(s, 5.9, y + 0.33, w=1.0, h=0.34)
    matrix(s, 7.3, y, m2, cw=0.56, ch=0.48, size=16, hl=hl)
    text(s, 9.3, y + 0.05, 1.8, 0.5, p, size=16, color=TEAL, bold=True)
    text(s, 9.3, y + 0.5, 1.8, 0.5, "in the diagonal", size=13, color=GREY)
    text(s, 11.2, y + 0.15, 1.5, 0.7, r, size=26, font=HEAD, bold=True)
text(s, 0.6, 6.55, 12, 0.4, "Full rank (2 pivots) means non-singular; a missing pivot means a zero row and a singular matrix.",
     size=15, color=GREY, italic=True)

# 37 · REF 3×3 -----------------------------------------------------------------
s = new_slide("Row echelon form of a 3×3 system", "Row echelon form")
eqs(s, 0.9, 1.75, ["a + b + 2c = 12", "3a − 3b − c = 3", "2a − b + 6c = 24"], size=20, gap=0.46, w=4.0, title="SYSTEM")
eqs(s, 8.3, 1.75, ["a + b + 2c = 12", "−6b − 7c = −33", "11c = 33"], size=20, gap=0.46, w=4.4, title="TRIANGULAR SYSTEM")
matrix(s, 1.2, 4.0, [["1", "1", "2"], ["3", "−3", "−1"], ["2", "−1", "6"]], cw=0.55, ch=0.46, size=16, label="Matrix")
arrow(s, 3.9, 4.6, w=3.9, h=0.36)
text(s, 3.6, 3.55, 4.5, 1.0, "R₂ ← R₂ − 3R₁        R₃ ← R₃ − 2R₁\nthen  R₃ ← 2R₃ − R₂", size=16, font=HEAD, color=NAVY, align="c")
matrix(s, 8.6, 4.0, [["1", "1", "2"], ["0", "−6", "−7"], ["0", "0", "11"]], cw=0.55, ch=0.46, size=16,
       hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, label="Row echelon form matrix")
text(s, 3.6, 5.3, 4.5, 1.2, "Three pivots: rank 3. The last row gives c = 3 at once, exactly as in the elimination we did by hand.",
     size=14, color=GREY, align="c", spacing=1.15)

# 38 · Normalizing the pivots -------------------------------------------------
s = new_slide("Normalizing the pivots to 1", "Row echelon form")
m_a = [["3", "*", "*", "*", "*"], ["0", "0", "1", "*", "*"], ["0", "0", "0", "−4", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
m_b = [["1", "*", "*", "*", "*"], ["0", "0", "1", "*", "*"], ["0", "0", "0", "1", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
hl = {(0, 0): PIV, (1, 2): PIV, (2, 3): PIV}
matrix(s, 0.9, 2.0, m_a, cw=0.5, ch=0.44, size=15, gap=0.04, hl=hl, label="Row echelon form, pivots 3, 1, −4")
for k, lab in enumerate(["÷ 3", "÷ 1", "÷ (−4)"]):
    y = 2.0 + k * 0.48
    arrow(s, 4.3, y + 0.07, w=1.3, h=0.3)
    text(s, 4.3, y - 0.3, 1.3, 0.35, lab, size=14, font=HEAD, color=NAVY, align="c")
matrix(s, 6.0, 2.0, m_b, cw=0.5, ch=0.44, size=15, gap=0.04, hl=hl, fill=RES, label="Row echelon form, every pivot 1")
defbox(s, 9.4, 2.0, 3.3, 3.0, [
    "Dividing a row by its pivot is a legal row operation (scaling by a non-zero constant).",
    "The staircase, the number of pivots and the rank are unchanged.",
], label="Why it is allowed", size=14)

# 39 · More examples -----------------------------------------------------------
s = new_slide("More examples", "Row echelon form")
ex = [
    ("Non-singular", [["1", "1", "1"], ["1", "2", "1"], ["1", "1", "2"]], "R₂ ← R₂ − R₁,  R₃ ← R₃ − R₁", None, [["1", "1", "1"], ["0", "1", "0"], ["0", "0", "1"]], {(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, "rank 3"),
    ("Singular", [["1", "1", "1"], ["1", "1", "2"], ["1", "1", "3"]], "R₂ ← R₂ − R₁,  R₃ ← R₃ − R₁", [["1", "1", "1"], ["0", "0", "1"], ["0", "0", "2"]], [["1", "1", "1"], ["0", "0", "1"], ["0", "0", "0"]], {(0, 0): PIV, (1, 2): PIV}, "rank 2"),
    ("Singular", [["1", "1", "1"], ["2", "2", "2"], ["3", "3", "3"]], "R₂ ← R₂ − 2R₁,  R₃ ← R₃ − 3R₁", None, [["1", "1", "1"], ["0", "0", "0"], ["0", "0", "0"]], {(0, 0): PIV}, "rank 1"),
]
for i, (t, m1, op, mid, m2, hl, r) in enumerate(ex):
    y = 1.65 + i * 1.72
    text(s, 0.6, y + 0.4, 2.0, 0.5, t, size=16, bold=True)
    matrix(s, 2.5, y, m1, cw=0.42, ch=0.4, size=14, gap=0.04)
    arrow(s, 4.4, y + 0.5, w=1.8, h=0.3, label=op, lsize=12, lw_=3.0)
    if mid:
        matrix(s, 6.5, y, mid, cw=0.42, ch=0.4, size=14, gap=0.04)
        arrow(s, 8.4, y + 0.5, w=1.2, h=0.3, label="R₃ ← R₃ − 2R₂", lsize=12, lw_=2.2)
        matrix(s, 9.9, y, m2, cw=0.42, ch=0.4, size=14, gap=0.04, hl=hl)
    else:
        matrix(s, 9.9, y, m2, cw=0.42, ch=0.4, size=14, gap=0.04, hl=hl)
    text(s, 11.75, y + 0.4, 1.2, 0.5, r, size=16, color=TEAL, bold=True)
text(s, 0.6, 6.75, 12, 0.3, "The second example needs two rounds: after the first round, row 3 is twice row 2.", size=12, color=GREY, italic=True)

# 40 · Pivots count the rank (3×3) -------------------------------------------
s = new_slide("Pivots count the rank", "Row echelon form")
pm = [
    ([["1", "1", "1"], ["1", "2", "1"], ["1", "1", "2"]], [["1", "1", "1"], ["0", "1", "0"], ["0", "0", "1"]], {(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, 3),
    ([["1", "1", "1"], ["1", "1", "2"], ["1", "1", "3"]], [["1", "1", "1"], ["0", "0", "1"], ["0", "0", "0"]], {(0, 0): PIV, (1, 2): PIV}, 2),
    ([["1", "1", "1"], ["2", "2", "2"], ["3", "3", "3"]], [["1", "1", "1"], ["0", "0", "0"], ["0", "0", "0"]], {(0, 0): PIV}, 1),
    ([["0", "0", "0"], ["0", "0", "0"], ["0", "0", "0"]], [["0", "0", "0"], ["0", "0", "0"], ["0", "0", "0"]], {}, 0),
]
text(s, 0.6, 1.6, 3, 0.3, "MATRIX", size=11, color=TEAL, bold=True)
text(s, 0.6, 3.75, 3, 0.3, "ROW ECHELON FORM", size=11, color=TEAL, bold=True)
for i, (m1, m2, hl, n) in enumerate(pm):
    x = 0.6 + i * 3.08
    text(s, x, 1.95, 2.9, 0.3, f"Matrix {i + 1}", size=14, bold=True)
    matrix(s, x + 0.55, 2.3, m1, cw=0.46, ch=0.4, size=14, gap=0.04)
    matrix(s, x + 0.55, 4.1, m2, cw=0.46, ch=0.4, size=14, gap=0.04, hl=hl)
    text(s, x, 5.45, 2.9, 0.4, f"Number of pivots = {n}", size=14, color=GREY, align="c")
    text(s, x, 5.85, 2.9, 0.6, f"Rank = {n}", size=24, font=HEAD, bold=True, align="c")

# ================================================= PART 5 · RREF ============
divider(5, "Reduced row echelon form",
        "Continue past the staircase: pivots become 1 and everything above a pivot becomes 0.")

# 42 · Definition RREF --------------------------------------------------------
s = new_slide("Reduced row echelon form: definition", "Reduced row echelon form")
defbox(s, 0.6, 1.55, 7.0, 3.4, [
    [("A matrix is in ", {}), ("reduced row echelon form", {"bold": True}), (" if", {})],
    "(1)  it is in row echelon form;",
    "(2)  every pivot equals 1;",
    "(3)  every pivot is the only non-zero entry in its column, so all entries above a pivot are 0.",
    [("Fact. ", {"bold": True}), ("Every matrix has exactly one reduced row echelon form; the rank is still the number of pivots.", {})],
])
text(s, 0.6, 5.15, 7.0, 1.4, "For a non-singular square matrix the reduced row echelon form is the identity matrix, "
     "and the augmented system is then solved outright.", size=15, color=GREY, spacing=1.15)
i5 = [["1", "0", "0", "0", "0"], ["0", "1", "0", "0", "0"], ["0", "0", "1", "0", "0"], ["0", "0", "0", "1", "0"], ["0", "0", "0", "0", "1"]]
matrix(s, 7.9, 1.7, i5, cw=0.4, ch=0.36, size=13, gap=0.04, hl=hl5, label="Rank 5", lsize=15, lcolor=NAVY)
r3 = [["1", "*", "0", "0", "*"], ["0", "0", "1", "0", "*"], ["0", "0", "0", "1", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
matrix(s, 10.5, 1.7, r3, cw=0.4, ch=0.36, size=13, gap=0.04, hl={(0, 0): PIV, (1, 2): PIV, (2, 3): PIV}, label="Rank 3", lsize=15, lcolor=NAVY)
text(s, 7.9, 4.5, 4.9, 2.0, [
    "•  Is in row echelon form.",
    "•  Each pivot is a 1.",
    "•  Any number above a pivot is 0.",
    "•  Rank of the matrix is the number of pivots.",
], size=14, color=GREY, spacing=1.15, after=4)

# 43 · RREF 2×2 ---------------------------------------------------------------
s = new_slide("Reduced row echelon form: a 2×2 example", "Reduced row echelon form")
matrix(s, 1.5, 2.3, [["1", "0.2"], ["0", "1"]], cw=0.65, ch=0.55, size=20, hl={(0, 1): PIV}, label="Row echelon form", lsize=15, lcolor=NAVY)
arrow(s, 4.0, 2.7, w=2.0, h=0.4, label="R₁ ← R₁ − 0.2 R₂", lsize=16, lw_=3.0)
matrix(s, 6.9, 2.3, [["1", "0"], ["0", "1"]], cw=0.65, ch=0.55, size=20, fill=RES, label="Reduced row echelon form", lsize=15, lcolor=NAVY)
rect(s, 0.6, 4.5, 12.1, 2.1, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 4.65, 11.5, 0.3, "THE COMPUTATION", size=12, color=TEAL, bold=True)
text(s, 0.9, 5.0, 11.5, 1.4, [
    "0.2 × [ 0   1 ]  =  [ 0   0.2 ]",
    "[ 1   0.2 ]  −  [ 0   0.2 ]  =  [ 1   0 ]       the entry above the second pivot is now 0",
], size=20, font=HEAD, color=NAVY, after=8)
text(s, 9.6, 2.2, 3.2, 1.4, "The only entry above a pivot was the 0.2. One row operation clears it.", size=15, color=GREY, spacing=1.15)

# 44 · General recipe ----------------------------------------------------------
s = new_slide("Reduced row echelon form: the general recipe", "Reduced row echelon form")
g1 = [["3", "*", "*", "*", "*"], ["0", "0", "2", "*", "*"], ["0", "0", "0", "−4", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
g2 = [["1", "*", "*", "*", "*"], ["0", "0", "1", "*", "*"], ["0", "0", "0", "1", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
g3 = [["1", "*", "0", "0", "*"], ["0", "0", "1", "0", "*"], ["0", "0", "0", "1", "*"], ["0", "0", "0", "0", "0"], ["0", "0", "0", "0", "0"]]
hl3 = {(0, 0): PIV, (1, 2): PIV, (2, 3): PIV}
matrix(s, 0.8, 2.1, g1, cw=0.46, ch=0.42, size=14, gap=0.04, hl=hl3, label="Row echelon form", lsize=14, lcolor=NAVY)
arrow(s, 3.75, 3.1, w=1.0, h=0.34, label="divide each row\nby its pivot", lsize=12, lw_=1.3)
matrix(s, 5.0, 2.1, g2, cw=0.46, ch=0.42, size=14, gap=0.04, hl=hl3)
arrow(s, 7.95, 3.1, w=1.0, h=0.34, label="clear entries\nabove pivots", lsize=12, lw_=1.3)
matrix(s, 9.2, 2.1, g3, cw=0.46, ch=0.42, size=14, gap=0.04, hl=hl3, fill=RES, label="Reduced row echelon form", lsize=14, lcolor=NAVY)
text(s, 0.6, 5.3, 12.1, 1.3, "Columns without a pivot keep their entries (the stars): they correspond to free variables of the system. "
     "The pivot columns become columns of the identity matrix.", size=16, color=GREY, spacing=1.2)

# 45 · RREF 3×3 example -------------------------------------------------------
s = new_slide("Reduced row echelon form: a 3×3 example", "Reduced row echelon form")
seq = [
    ([["1", "2", "3"], ["0", "1", "4"], ["0", "0", "1"]], "Row echelon form", "R₁ ← R₁ − 2R₂"),
    ([["1", "0", "−5"], ["0", "1", "4"], ["0", "0", "1"]], None, "R₁ ← R₁ + 5R₃"),
    ([["1", "0", "0"], ["0", "1", "4"], ["0", "0", "1"]], None, "R₂ ← R₂ − 4R₃"),
    ([["1", "0", "0"], ["0", "1", "0"], ["0", "0", "1"]], "Reduced row echelon form", None),
]
for i, (m, lab, op) in enumerate(seq):
    x = 0.7 + i * 3.25
    matrix(s, x, 2.3, m, cw=0.5, ch=0.46, size=16, gap=0.04, fill=RES if i == 3 else CELL,
           hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, label=lab, lsize=14, lcolor=NAVY)
    if op:
        arrow(s, x + 2.1, 2.95, w=1.0, h=0.34, label=op, lsize=14, lw_=2.0)
text(s, 0.6, 5.0, 12.1, 1.5, "Work from the last pivot upwards, or from the first pivot downwards: either order clears every entry above a pivot "
     "and never disturbs a column already finished. The result is the identity, so the matrix has rank 3 and is non-singular.",
     size=16, color=GREY, spacing=1.2)

# ================================================ PART 6 · GAUSSIAN ELIM ====
divider(6, "The Gaussian elimination algorithm",
        "Putting it all together on the augmented matrix: pivot, reduce, back-substitute, and read the answer.")

# 47 · Augmented matrix -------------------------------------------------------
s = new_slide("The augmented matrix", "Gaussian elimination")
defbox(s, 0.6, 1.55, 12.1, 1.55, [
    [("The ", {}), ("augmented matrix", {"bold": True}), (" [ A | b ] of the system A·x = b is the coefficient matrix A with the column "
     "of constants b appended on the right. Row operations act on entire rows, constants included, so the system stays equivalent.", {})],
])
eqs(s, 0.9, 3.5, ["2a − b + c = 1", "2a + 2b + 4c = −2", "4a + b = 4"], size=24, gap=0.62, w=4.5, title="SYSTEM")
arrow(s, 5.7, 4.35, w=1.2, h=0.38)
text(s, 7.3, 3.5, 5, 0.35, "AUGMENTED MATRIX", size=14, color=GREY, bold=True)
text(s, 7.3, 3.95, 0.5, 1.9, "R₁\nR₂\nR₃", size=18, font=HEAD, color=GREY, spacing=1.0)
matrix(s, 7.8, 3.95, [["2", "−1", "1", "1"], ["2", "2", "4", "−2"], ["4", "1", "0", "4"]], cw=0.6, ch=0.55, size=18, aug=2)
text(s, 0.9, 5.75, 11.6, 0.8, "Each row is one equation; the bar separates the coefficients of a, b, c from the constant. "
     "Now proceed with the elimination method on the rows.", size=16, color=GREY, spacing=1.2)

# 48 · Pivoting on the first column ------------------------------------------
s = new_slide("Pivoting on the first column", "Gaussian elimination")
matrix(s, 0.8, 1.85, [["2", "−1", "1", "1"], ["2", "2", "4", "−2"], ["4", "1", "0", "4"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV}, label="Start: pivot 2 in row 1", lsize=14, lcolor=NAVY)
arrow(s, 4.1, 2.5, w=0.9, h=0.34)
matrix(s, 5.3, 1.85, [["1", "−½", "½", "½"], ["2", "2", "4", "−2"], ["4", "1", "0", "4"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV}, label="R₁ ← ½ R₁", lsize=14, lcolor=NAVY)
arrow(s, 8.6, 2.5, w=0.9, h=0.34)
matrix(s, 9.8, 1.85, [["1", "−½", "½", "½"], ["0", "3", "3", "−3"], ["0", "3", "−2", "2"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV, (1, 0): RES, (2, 0): RES}, label="R₂ ← R₂ − 2R₁,   R₃ ← R₃ − 4R₁", lsize=14, lcolor=NAVY)
rect(s, 0.6, 4.55, 12.1, 2.0, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 4.7, 11.5, 0.3, "THE ROW ARITHMETIC", size=12, color=TEAL, bold=True)
text(s, 0.9, 5.05, 11.5, 1.4, [
    "R₂ − 2R₁ :   [ 2   2   4 | −2 ]  −  2 · [ 1   −½   ½ | ½ ]   =   [ 0   3   3 | −3 ]",
    "R₃ − 4R₁ :   [ 4   1   0 | 4 ]  −  4 · [ 1   −½   ½ | ½ ]   =   [ 0   3   −2 | 2 ]",
], size=19, font=HEAD, color=NAVY, after=8)

# 49 · Pivoting on columns 2 and 3 -------------------------------------------
s = new_slide("Pivoting on the second and third columns", "Gaussian elimination")
matrix(s, 0.8, 1.85, [["1", "−½", "½", "½"], ["0", "3", "3", "−3"], ["0", "3", "−2", "2"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(1, 1): PIV}, label="R₂ ← ⅓ R₂", lsize=14, lcolor=NAVY)
arrow(s, 4.1, 2.5, w=0.9, h=0.34)
matrix(s, 5.3, 1.85, [["1", "−½", "½", "½"], ["0", "1", "1", "−1"], ["0", "0", "−5", "5"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(1, 1): PIV, (2, 1): RES, (2, 2): PIV}, label="R₃ ← R₃ − 3R₂,   then  R₃ ← −⅕ R₃", lsize=14, lcolor=NAVY)
arrow(s, 8.6, 2.5, w=0.9, h=0.34)
matrix(s, 9.8, 1.85, [["1", "−½", "½", "½"], ["0", "1", "1", "−1"], ["0", "0", "1", "−1"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, fill=RES, label="Row echelon form", lsize=14, lcolor=NAVY)
rect(s, 0.6, 4.55, 12.1, 2.0, fill=TINT, rounded=True, radius=0.05)
text(s, 0.9, 4.7, 5.5, 0.3, "THE ROW ARITHMETIC", size=12, color=TEAL, bold=True)
text(s, 0.9, 5.05, 6.5, 1.4, [
    "R₃ − 3R₂ :  [ 0  3  −2 | 2 ] − 3 · [ 0  1  1 | −1 ]  =  [ 0  0  −5 | 5 ]",
    "−⅕ R₃ :  [ 0  0  −5 | 5 ] ÷ (−5)  =  [ 0  0  1 | −1 ]",
], size=17, font=HEAD, color=NAVY, after=8)
text(s, 7.9, 4.7, 4.5, 0.3, "AS EQUATIONS", size=12, color=TEAL, bold=True)
eqs(s, 7.9, 5.0, ["a − ½ b + ½ c = ½", "b + c = −1", "c = −1"], size=19, gap=0.45, w=4.5)

# 50 · Back substitution ------------------------------------------------------
s = new_slide("Back substitution", "Gaussian elimination")
matrix(s, 0.8, 1.85, [["1", "−½", "½", "½"], ["0", "1", "1", "−1"], ["0", "0", "1", "−1"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV}, label="Row echelon form", lsize=14, lcolor=NAVY)
arrow(s, 4.1, 2.5, w=0.9, h=0.34)
matrix(s, 5.3, 1.85, [["1", "−½", "0", "1"], ["0", "1", "0", "0"], ["0", "0", "1", "−1"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV, (0, 2): RES, (1, 2): RES}, label="R₂ ← R₂ − R₃,   R₁ ← R₁ − ½ R₃", lsize=14, lcolor=NAVY)
arrow(s, 8.6, 2.5, w=0.9, h=0.34)
matrix(s, 9.8, 1.85, [["1", "0", "0", "1"], ["0", "1", "0", "0"], ["0", "0", "1", "−1"]], cw=0.6, ch=0.52, size=17, aug=2,
       hl={(0, 0): PIV, (1, 1): PIV, (2, 2): PIV, (0, 1): RES}, label="R₁ ← R₁ + ½ R₂", lsize=14, lcolor=NAVY)
rect(s, 0.6, 4.55, 6.3, 2.0, fill=RES, rounded=True, radius=0.05)
text(s, 0.9, 4.7, 5.8, 0.3, "THE RESULT", size=12, color=TEAL, bold=True)
text(s, 0.9, 5.05, 5.8, 0.6, "a = 1,   b = 0,   c = −1", size=28, font=HEAD, bold=True)
text(s, 0.9, 5.75, 5.8, 0.6, "The coefficient part is the identity matrix; the constants column is the solution.", size=14, color=GREY)
check(s, 7.3, 4.65, ["2(1) − 0 + (−1) = 1  ✓", "2(1) + 2(0) + 4(−1) = −2  ✓", "4(1) + 0 = 4  ✓"], size=16)

# 51 · Singular case in Gaussian elimination ---------------------------------
s = new_slide("What if the system is singular?", "Gaussian elimination")
text(s, 0.6, 1.55, 12, 0.4, "Look at the row of zeros, then at its constant.", size=17, color=GREY)
matrix(s, 0.8, 2.1, [["1", "2", "−1", "5"], ["2", "4", "5", "1"], ["3", "6", "4", "6"]], cw=0.56, ch=0.48, size=16, aug=2)
arrow(s, 4.1, 2.7, w=1.3, h=0.34, label="after row reduction", lsize=12, lw_=2.0)
matrix(s, 5.8, 2.1, [["1", "2", "−1", "5"], ["0", "0", "7", "−9"], ["0", "0", "0", "0"]], cw=0.56, ch=0.48, size=16, aug=2,
       hl={(2, 0): PIV, (2, 1): PIV, (2, 2): PIV, (2, 3): RES})
text(s, 9.3, 2.1, 3.5, 1.6, [[("0a + 0b + 0c = 0", {"font": HEAD, "size": 18})], [("Infinitely many solutions", {"bold": True, "color": TEAL, "size": 18})],
                              "the constant in the zero row is 0"], size=14, color=GREY, after=4)
matrix(s, 0.8, 4.2, [["1", "2", "−1", "5"], ["2", "4", "5", "1"], ["3", "6", "4", "10"]], cw=0.56, ch=0.48, size=16, aug=2, hl={(2, 3): PIV})
arrow(s, 4.1, 4.8, w=1.3, h=0.34, label="after row reduction", lsize=12, lw_=2.0)
matrix(s, 5.8, 4.2, [["1", "2", "−1", "5"], ["0", "0", "7", "−9"], ["0", "0", "0", "4"]], cw=0.56, ch=0.48, size=16, aug=2,
       hl={(2, 0): PIV, (2, 1): PIV, (2, 2): PIV, (2, 3): "F6D5D2"})
text(s, 9.3, 4.2, 3.5, 1.6, [[("0a + 0b + 0c = 4", {"font": HEAD, "size": 18})], [("No solutions", {"bold": True, "color": RED, "size": 18})],
                              "the constant in the zero row is not 0"], size=14, color=GREY, after=4)
text(s, 0.6, 6.2, 12, 0.7, "Row operations used: R₂ ← R₂ − 2R₁,  R₃ ← R₃ − 3R₁,  R₃ ← R₃ − R₂.  Only the constants column differs between the two systems.",
     size=14, color=GREY, italic=True)

# 52 · Summary algorithm -------------------------------------------------------
s = new_slide("Gaussian elimination: the algorithm", "Gaussian elimination")
steps = [
    ("Create the augmented matrix", "Write [ A | b ]: one row per equation, the constants in the last column."),
    ("Reach row echelon form", "Pivot column by column: scale the pivot row to make the pivot 1, then subtract multiples of it from the rows below."),
    ("Reduce further or back-substitute", "Clear the entries above each pivot to reach reduced row echelon form, or solve from the last equation upwards."),
    ("Stop at a row of zeros", "A row  0 … 0 | 0  means infinitely many solutions; a row  0 … 0 | β  with β ≠ 0 means no solution. "
                               "Otherwise the solution is unique: the last column."),
]
for i, (t, b) in enumerate(steps):
    x = 0.6 + i * 3.08
    card(s, x, 1.7, 2.9, 3.55, t, b, number=i + 1, tsize=16, bsize=14)
    if i < 3:
        arrow(s, x + 2.88, 3.5, w=0.26, h=0.3)
defbox(s, 0.6, 5.45, 12.1, 1.35, [
    [("Gaussian elimination", {"bold": True}), (" is this procedure: elementary row operations applied to the augmented matrix until it is in (reduced) row echelon form.", {})],
], label="Definition", size=15)

# 53 · Recap -------------------------------------------------------------------
s = new_slide(dark=True)
text(s, 0.6, 0.5, 8, 0.4, "WEEK 2 RECAP", size=14, color=AMBER_L, bold=True)
text(s, 0.6, 0.9, 11.5, 0.8, "The definitions to remember", size=36, font=HEAD, color=WHITE, bold=True)
recap = [
    ("Linear system", "A finite set of linear equations in the same variables; a solution satisfies all of them at once."),
    ("Equivalent systems", "Same solutions. Swapping, scaling by k ≠ 0, and adding a multiple of one equation to another keep a system equivalent."),
    ("Singular vs non-singular", "Non-singular: exactly one solution. Singular: infinitely many (redundant) or none (contradictory)."),
    ("Rank", "The number of independent pieces of information; number of variables minus the dimension of the solution space; the number of pivots."),
    ("Row echelon form", "Zero rows at the bottom; each pivot strictly to the right of the pivot above it."),
    ("Reduced row echelon form", "Row echelon form with every pivot equal to 1 and zeros above every pivot. Unique for each matrix."),
    ("Augmented matrix", "[ A | b ]: the coefficient matrix with the constants column attached."),
    ("Gaussian elimination", "Row-reduce the augmented matrix to (reduced) row echelon form; read the solution or detect a zero row."),
]
for i, (t, b) in enumerate(recap):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.85 + row * 1.25
    rect(s, x, y, 5.95, 1.1, fill="1D2D52", rounded=True, radius=0.05)
    text(s, x + 0.25, y + 0.12, 5.5, 0.35, t, size=16, bold=True, color=AMBER_L)
    text(s, x + 0.25, y + 0.45, 5.5, 0.65, b, size=12, color=LIGHT, spacing=1.1)

prs.save("Linear_Algebra_Week2_Slides.pptx")
print("saved", len(prs.slides), "slides")
